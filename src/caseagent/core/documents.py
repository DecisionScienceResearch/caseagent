"""Turn evidence files into text the model can read.

Supports the file types the DecisionLayer forms accept (.pdf, .docx, .rtf, .txt)
plus other plain-text formats. Anything else is reported as unreadable rather
than guessed at, so the agent never argues from a document it could not see.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

PLAIN_TEXT_SUFFIXES = {".txt", ".md", ".csv", ".tsv", ".json", ".log", ".eml", ".html", ".htm"}
SUPPORTED_SUFFIXES = PLAIN_TEXT_SUFFIXES | {".pdf", ".docx", ".rtf"}

# pypdf logs recoverable structural problems ("Ignoring wrong pointing object")
# at WARNING for PDFs that still extract correctly.
logging.getLogger("pypdf").setLevel(logging.ERROR)


@dataclass(frozen=True)
class ExtractedText:
    text: str
    warning: str | None = None

    @property
    def readable(self) -> bool:
        return bool(self.text.strip())


def normalize_text(text: str) -> str:
    """Normalize newlines and trim trailing whitespace without touching content."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [line.rstrip() for line in text.split("\n")]
    return "\n".join(lines).strip()


def extract_text(path: Path) -> ExtractedText:
    """Extract text from one file. Never raises for an unreadable document."""
    suffix = path.suffix.lower()
    try:
        if suffix in PLAIN_TEXT_SUFFIXES:
            return ExtractedText(normalize_text(path.read_text(encoding="utf-8", errors="replace")))
        if suffix == ".pdf":
            return _extract_pdf(path)
        if suffix == ".docx":
            return _extract_docx(path)
        if suffix == ".rtf":
            return _extract_rtf(path)
    except Exception as exc:  # a corrupt file should not sink the whole case
        return ExtractedText("", warning=f"could not read {path.name}: {exc}")
    return ExtractedText(
        "", warning=f"{path.name}: unsupported file type {suffix or '(none)'}; content not read"
    )


def _extract_pdf(path: Path) -> ExtractedText:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [normalize_text(page.extract_text() or "") for page in reader.pages]
    if not any(pages):
        return ExtractedText(
            "",
            warning=f"{path.name}: no extractable text (scanned PDF?). Run OCR and re-add it.",
        )
    if len(pages) == 1:
        return ExtractedText(pages[0])
    marked = [f"[Page {number}]\n{text}" for number, text in enumerate(pages, start=1)]
    return ExtractedText("\n\n".join(marked))


def _extract_docx(path: Path) -> ExtractedText:
    import docx

    document = docx.Document(str(path))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ExtractedText(normalize_text("\n".join(blocks)))


def _extract_rtf(path: Path) -> ExtractedText:
    from striprtf.striprtf import rtf_to_text

    raw = path.read_text(encoding="utf-8", errors="replace")
    return ExtractedText(normalize_text(rtf_to_text(raw)))
