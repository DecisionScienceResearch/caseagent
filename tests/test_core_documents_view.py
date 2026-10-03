from pathlib import Path

import docx
from conftest import human_filing, make_pdf, seed_thread

from caseagent.core.casefile import save_case
from caseagent.core.documents import extract_text
from caseagent.core.models import Filing, Role, Stage, Statement
from caseagent.core.view import build_stage_view, exhibit_label


def test_extract_pdf_docx_rtf_txt(tmp_path: Path):
    (tmp_path / "a.pdf").write_bytes(make_pdf("Hello exhibit"))
    assert "Hello exhibit" in extract_text(tmp_path / "a.pdf").text

    document = docx.Document()
    document.add_paragraph("Clause one.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Rent"
    table.rows[0].cells[1].text = "$20,000"
    document.save(tmp_path / "b.docx")
    text = extract_text(tmp_path / "b.docx").text
    assert "Clause one." in text and "Rent | $20,000" in text

    (tmp_path / "c.rtf").write_text(r"{\rtf1\ansi Plain {\b bold} text.}")
    assert "Plain bold text." in extract_text(tmp_path / "c.rtf").text

    (tmp_path / "d.txt").write_text("line one\r\nline two  \r\n")
    assert extract_text(tmp_path / "d.txt").text == "line one\nline two"


def test_unreadable_files_are_reported_not_guessed(tmp_path: Path):
    (tmp_path / "scan.png").write_bytes(b"\x89PNG")
    result = extract_text(tmp_path / "scan.png")
    assert result.text == "" and "unsupported" in result.warning

    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    result = extract_text(tmp_path / "broken.pdf")
    assert result.text == "" and result.warning


def test_exhibit_label():
    assert exhibit_label("Exhibit A-2 - NY Liquor License Issuance.pdf") == "Exhibit A-2"
    assert exhibit_label("Exhibit C-6B - Driver Account History.pdf") == "Exhibit C-6B"
    assert exhibit_label("invoice.pdf") == "invoice"


def test_answer_view_sees_claim_and_own_exhibits_only(case_bundle):
    store = seed_thread(case_bundle, Stage.CLAIM)
    view = build_stage_view(case_bundle, Stage.ANSWER, store.thread())

    assert view.party.name == "Arcology LLC"
    assert view.opponent.name == "BevMart Inc."
    assert [f.stage for f in view.record] == [Stage.CLAIM]
    assert view.record[0].argument.startswith("HUMAN CLAIM")
    names = [exhibit.name for exhibit in view.exhibits]
    assert names == [
        "Exhibit A-1 - POS Ledger.txt",
        "Exhibit A-2 - License.txt",
        "Exhibit R-1 - Notice.txt",
    ]
    by_name = {exhibit.name: exhibit for exhibit in view.exhibits}
    assert by_name["Exhibit A-1 - POS Ledger.txt"].is_new is False
    assert by_name["Exhibit R-1 - Notice.txt"].is_new is True
    assert "Commencement Date established" in by_name["Exhibit R-1 - Notice.txt"].text
    # Later-phase exhibits are invisible.
    assert "Exhibit A-3 - Emails.txt" not in names
    assert "Exhibit R-3 - LOI.txt" not in names


def test_claim_view_has_no_record(case_bundle):
    view = build_stage_view(case_bundle, Stage.CLAIM, {})
    assert view.record == ()
    assert [e.name for e in view.exhibits] == [
        "Exhibit A-1 - POS Ledger.txt",
        "Exhibit A-2 - License.txt",
    ]
    assert all(exhibit.is_new for exhibit in view.exhibits)
    assert view.pending_demands is None


def test_skipped_phase_contributes_nothing(case_bundle):
    # A Reply written against a Claim alone: no Answer filing, no R-1.
    thread = {Stage.CLAIM: human_filing(Stage.CLAIM)}
    view = build_stage_view(case_bundle, Stage.REPLY, thread)
    assert [f.stage for f in view.record] == [Stage.CLAIM]
    assert "Exhibit R-1 - Notice.txt" not in [e.name for e in view.exhibits]
    assert "Exhibit A-3 - Emails.txt" in [e.name for e in view.exhibits]


def test_pending_demands_and_counterclaim_flags(case_bundle):
    thread = {
        Stage.CLAIM: Filing(stage=Stage.CLAIM, argument="c", evidence_demands="the lease file"),
        Stage.ANSWER: Filing(
            stage=Stage.ANSWER,
            argument="a",
            counterclaim={"argument": "owed", "amount_usd": "42000"},
            evidence_demands="the POS data",
        ),
    }
    answer_view = build_stage_view(case_bundle, Stage.ANSWER, {Stage.CLAIM: thread[Stage.CLAIM]})
    assert answer_view.pending_demands == "the lease file"
    reply_view = build_stage_view(case_bundle, Stage.REPLY, thread)
    assert reply_view.pending_demands == "the POS data"
    assert reply_view.counterclaim_filed is True


def test_statements_are_private_to_their_side(case_bundle):
    case = case_bundle.case.model_copy(
        update={
            "statements": {
                Role.CLAIMANT: [Statement(question="Opened?", answer="July 1")],
                Role.RESPONDENT: [Statement(question="Notice sent?", answer="May 1")],
            }
        }
    )
    bundle = save_case(case, case_bundle.root)
    claim_view = build_stage_view(bundle, Stage.CLAIM, {})
    assert [s.answer for s in claim_view.statements] == ["July 1"]
    answer_view = build_stage_view(bundle, Stage.ANSWER, {Stage.CLAIM: human_filing(Stage.CLAIM)})
    assert [s.answer for s in answer_view.statements] == ["May 1"]
